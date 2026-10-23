"""Gold workbook (v0 layout) -> ``GoldRecord`` (FormulationRecord + originals + flags).

Lenient by design: human/gold labels are never rejected for breaking a guide
rule. Each record keeps

- ``original``: the raw rows from Records / Components / Results, unchanged
- ``record``: a v3 ``FormulationRecord`` (None only if it cannot be built)
- ``flags``: every normalization, ambiguity, invalid value and rule violation
- ``gold_record_id``: the spreadsheet's record ID, preserved as is; extracted
  records are matched to gold explicitly (``negative_lines.gold.match``), our
  canonical record-ID generation is not changed to imitate the spreadsheet.

Paths are configurable; by default only the committed development workbook is
read. Loading benchmark documents requires ``allow_sets=("benchmark",)``.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from negative_lines.gold import vocab_map
from negative_lines.gold.split import BenchmarkAccessError, Split, get_split

__all__ = ["BenchmarkAccessError", "GoldFlag", "GoldRecord", "GoldSet", "load_gold"]
from negative_lines.schema.enums import VOCAB_VERSION
from negative_lines.schema.record import (
    Component,
    Developer,
    Exposure,
    FormulationRecord,
    Polymer,
    Process,
    Provenance,
    Result,
    SoftbakeStep,
)
from negative_lines.vocab import get_vocab


class GoldFlag(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sheet: str
    row: int  # 1-based spreadsheet row
    field: str
    original: Any = None
    value: Any = None
    status: vocab_map.FlagStatus
    note: str


class GoldRecord(BaseModel):
    model_config = ConfigDict(extra="forbid", arbitrary_types_allowed=True)

    gold_record_id: str
    doc_id: str
    example_label: str
    set: str | None
    labeler: str | None
    labeler_confidence: str | None
    low_confidence: bool
    record: FormulationRecord | None
    flags: list[GoldFlag] = Field(default_factory=list)
    load_errors: list[str] = Field(default_factory=list)
    original: dict[str, Any] = Field(default_factory=dict)
    reconciliation: list[dict[str, Any]] = Field(default_factory=list)


class GoldSet(BaseModel):
    model_config = ConfigDict(extra="forbid")

    path: str
    sha256: str
    gold_version: str
    documents: list[dict[str, Any]]
    records: list[GoldRecord]

    def for_doc(self, doc_id: str) -> list[GoldRecord]:
        return [r for r in self.records if r.doc_id == doc_id]

    def get(self, gold_record_id: str) -> GoldRecord:
        for r in self.records:
            if r.gold_record_id == gold_record_id:
                return r
        raise KeyError(gold_record_id)

    def flag_summary(self) -> Counter:
        return Counter((f.status, f.sheet, f.field) for r in self.records for f in r.flags)


# ----- low-level parsing --------------------------------------------------

_NUM = re.compile(r"^\s*-?\d+(?:\.\d+)?\s*$")


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _str(v: Any) -> str | None:
    if _blank(v):
        return None
    return str(v).strip()


class _Ctx:
    """Collects flags for one record."""

    def __init__(self) -> None:
        self.flags: list[GoldFlag] = []

    def flag(self, sheet: str, row: int, field: str, original: Any, value: Any, status: str,
             note: str) -> None:  # fmt: skip
        self.flags.append(GoldFlag(sheet=sheet, row=row, field=field, original=original,
                                   value=value, status=status, note=note))  # fmt: skip

    def number(self, sheet: str, row: int, field: str, v: Any) -> float | None:
        if _blank(v):
            return None
        if isinstance(v, bool):
            self.flag(sheet, row, field, v, None, "invalid", "boolean where a number is expected")
            return None
        if isinstance(v, int | float):
            return float(v)
        if isinstance(v, str) and _NUM.match(v):
            return float(v)
        self.flag(sheet, row, field, v, None, "invalid", "not a single number; original kept")
        return None

    def number_or_text(self, v: Any) -> float | str | None:
        if _blank(v):
            return None
        if isinstance(v, int | float) and not isinstance(v, bool):
            return float(v)
        s = str(v).strip()
        return float(s) if _NUM.match(s) else s

    def enum(self, sheet: str, row: int, field: str, v: Any, allowed: set[str]) -> str | None:
        s = _str(v)
        if s is None:
            return None
        if s in allowed:
            return s
        self.flag(sheet, row, field, v, None, "invalid", f"not in vocab ({len(allowed)} values)")
        return None


def _rows(ws) -> list[tuple[int, dict[str, Any]]]:
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(it)]
    out = []
    for i, r in enumerate(it, start=2):
        if r is None or all(_blank(v) for v in r):
            continue
        out.append((i, dict(zip(header, r, strict=False))))
    return out


def _bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"1", "true", "yes", "y"}


# ----- record construction -----------------------------------------------


def _primer_or_underlayer(ctx: _Ctx, row: int, v: Any) -> tuple[str | None, str | None]:
    s = _str(v)
    if s is None:
        return None, None
    t = s.casefold()
    if re.search(r"underlayer|under layer|lower (organic )?layer|bottom layer", t):
        ctx.flag("Records", row, "primer_or_underlayer", s, "underlayer", "normalized",
                 "merged column -> process.underlayer")  # fmt: skip
        return None, s
    if re.search(r"hmds|primer|omnicoat|adhesion promoter|hexamethyldisilazane", t):
        ctx.flag("Records", row, "primer_or_underlayer", s, "primer", "normalized",
                 "merged column -> process.primer")  # fmt: skip
        return s, None
    ctx.flag("Records", row, "primer_or_underlayer", s, "primer", "ambiguous",
             "neither primer nor underlayer recognized; stored in process.primer")  # fmt: skip
    return s, None


def _component(ctx: _Ctx, row: int, c: dict[str, Any]) -> Component | None:
    vocab = get_vocab()
    raw = ctx.number_or_text(c.get("amount_raw"))
    mixture_id = mixture_ratio = None
    if isinstance(raw, str):
        mix = vocab_map.parse_mixture(raw)
        if mix:
            mixture_id, mixture_ratio = mix
            ctx.flag("Components", row, "amount_raw", raw, {"mixture_id": mixture_id,
                     "mixture_ratio": mixture_ratio}, "normalized",
                     "mixture total parsed into mixture_id/mixture_ratio; amount_raw kept")  # fmt: skip
    basis = ctx.enum("Components", row, "amount_basis", c.get("amount_basis"),
                     set(vocab.enum("amount_bases")))  # fmt: skip
    conf = ctx.enum("Components", row, "labeler_confidence", c.get("labeler_confidence"),
                    set(vocab.enum("labeler_confidence")))  # fmt: skip
    if basis == "unspecified" and conf != "low":
        ctx.flag("Components", row, "labeler_confidence", conf, conf, "rule_violation",
                 "amount_basis 'unspecified' requires labeler_confidence 'low' (vocab)")  # fmt: skip
    polymer = None
    poly = {k: c.get(f"polymer_{k}") for k in ("repeat_units", "monomer_ratio", "mw", "pdi")}
    if any(not _blank(v) for v in poly.values()):
        polymer = Polymer(
            repeat_units=_str(poly["repeat_units"]),
            monomer_ratio=_str(poly["monomer_ratio"]),
            mw=ctx.number("Components", row, "polymer_mw", poly["mw"]),
            pdi=ctx.number("Components", row, "polymer_pdi", poly["pdi"]),
        )
    role = ctx.enum("Components", row, "role", c.get("role"), set(vocab.enum("roles")))
    try:
        return Component(
            role=role or "other",
            role_other_text=_str(c.get("role_other_text"))
            or (None if role else "(invalid role in gold)"),
            name_as_written=_str(c.get("name_as_written")) or "(blank)",
            definition_location=_str(c.get("definition_location")),
            polymer=polymer,
            amount_raw=raw,
            amount_unit=_str(c.get("amount_unit")),
            amount_basis=basis,
            mixture_id=mixture_id,
            mixture_ratio=mixture_ratio,
            labeler_confidence=conf,
            labeler_note=_str(c.get("labeler_note")),
        )
    except ValidationError as exc:
        ctx.flag("Components", row, "<component>", None, None, "invalid", f"not loadable: {exc}")
        return None


def _result(ctx: _Ctx, row: int, s: dict[str, Any]) -> Result | None:
    vocab = get_vocab()
    metric = ctx.enum("Results", row, "metric", s.get("metric"), set(vocab.enum("metrics")))
    written = _str(s.get("metric_as_written"))
    if metric is not None:
        metric, f = vocab_map.normalize_metric(metric, written)
        if f:
            ctx.flag("Results", row, f.field, f.original, f.value, f.status, f.note)
    sigma = ctx.number("Results", row, "sigma", s.get("sigma"))
    film = ctx.number("Results", row, "film_thickness", s.get("film_thickness"))
    if film is not None:
        ctx.flag("Results", row, "film_thickness", s.get("film_thickness"), film, "ambiguous",
                 "unit not annotated in gold v0; record field is nm, value kept as is")  # fmt: skip
    try:
        return Result(
            metric=metric or "other",
            metric_as_written=written or "(blank)",
            value=ctx.number("Results", row, "value", s.get("value")),
            value_qualitative=_str(s.get("value_qualitative")),
            scale=_str(s.get("scale")),
            unit=_str(s.get("unit")),
            sigma=int(sigma) if sigma is not None else None,
            pitch_nm=ctx.number("Results", row, "pitch_nm", s.get("pitch_nm")),
            pattern_type=ctx.enum(
                "Results",
                row,
                "pattern_type",
                s.get("pattern_type"),
                set(vocab.enum("pattern_types")),
            ),  # fmt: skip
            film_thickness=film,
            condition=_str(s.get("condition")),
            method=_str(s.get("method")),
            reference_record_id=_str(s.get("reference_record_id")),
            labeler_confidence=ctx.enum(
                "Results",
                row,
                "labeler_confidence",
                s.get("labeler_confidence"),
                set(vocab.enum("labeler_confidence")),
            ),  # fmt: skip
            labeler_note=_str(s.get("labeler_note")),
        )
    except ValidationError as exc:
        ctx.flag("Results", row, "<result>", None, None, "invalid", f"not loadable: {exc}")
        return None


def _time_unit(ctx: _Ctx, row: int, field: str, v: Any) -> str | None:
    return ctx.enum("Records", row, field, v, set(get_vocab().time_units))


def _process(ctx: _Ctx, row: int, r: dict[str, Any]) -> Process:
    vocab = get_vocab()
    primer, underlayer = _primer_or_underlayer(ctx, row, r.get("primer_or_underlayer"))

    source = ctx.enum("Records", row, "exposure_source", r.get("exposure_source"),
                      set(vocab.enum("exposure_sources")))  # fmt: skip
    source, f = vocab_map.normalize_exposure_source(source, _str(r.get("exposure_tool")))
    if f:
        ctx.flag("Records", row, f.field, f.original, f.value, f.status, f.note)
    unit, f = vocab_map.normalize_dose_unit(_str(r.get("dose_unit")))
    if f:
        ctx.flag("Records", row, f.field, f.original, f.value, f.status, f.note)
    dose = ctx.number("Records", row, "dose", r.get("dose"))
    if dose is not None and unit is None:
        ctx.flag("Records", row, "dose", r.get("dose"), None, "ambiguous",
                 "dose without a valid unit left unset; original kept")  # fmt: skip
        dose = None

    steps = []
    sb_temp, sb_time = r.get("softbake_temp"), r.get("softbake_time")
    if not (_blank(sb_temp) and _blank(sb_time)):
        steps.append(
            SoftbakeStep(
                temp=ctx.number_or_text(sb_temp),
                time=ctx.number("Records", row, "softbake_time", sb_time),
                time_unit=_time_unit(ctx, row, "softbake_time_unit", r.get("softbake_time_unit")),
            )
        )

    exposure = Exposure(
        source=source,
        mode=ctx.enum(
            "Records",
            row,
            "exposure_mode",
            r.get("exposure_mode"),
            set(vocab.enum("exposure_modes")),
        ),  # fmt: skip
        tool=_str(r.get("exposure_tool")),
        dose=dose,
        dose_unit=unit,
        dose_type=ctx.enum(
            "Records", row, "dose_type", r.get("dose_type"), set(vocab.enum("dose_types"))
        ),  # fmt: skip
        intensity_mw_cm2=ctx.number("Records", row, "intensity_mw_cm2", r.get("intensity_mw_cm2")),
        exposure_time_s=ctx.number("Records", row, "exposure_time_s", r.get("exposure_time_s")),
        ebeam_kv=ctx.number("Records", row, "ebeam_kv", r.get("ebeam_kv")),
    )
    return Process(
        substrate=_str(r.get("substrate")),
        primer=primer,
        underlayer=underlayer,
        coating_method=_str(r.get("coating_method")),
        film_thickness=ctx.number("Records", row, "film_thickness", r.get("film_thickness")),
        film_thickness_unit=ctx.enum(
            "Records",
            row,
            "film_thickness_unit",
            r.get("film_thickness_unit"),
            set(vocab.length_units),
        ),  # fmt: skip
        thickness_stage=ctx.enum(
            "Records",
            row,
            "thickness_stage",
            r.get("thickness_stage"),
            set(vocab.enum("thickness_stages")),
        ),  # fmt: skip
        softbake_steps=steps,
        exposure=exposure,
        peb_temp=ctx.number_or_text(r.get("peb_temp")),
        peb_time=ctx.number("Records", row, "peb_time", r.get("peb_time")),
        peb_time_unit=_time_unit(ctx, row, "peb_time_unit", r.get("peb_time_unit")),
        developer=Developer(
            name=_str(r.get("developer_name")),
            concentration=_str(r.get("developer_concentration")),
            time=ctx.number("Records", row, "developer_time", r.get("developer_time")),
            time_unit=_time_unit(ctx, row, "developer_time_unit", r.get("developer_time_unit")),
            method=ctx.enum(
                "Records",
                row,
                "developer_method",
                r.get("developer_method"),
                set(vocab.enum("developer_methods")),
            ),  # fmt: skip
        ),
        rinse=_str(r.get("rinse")),
        hardbake_temp=ctx.number_or_text(r.get("hardbake_temp")),
        hardbake_time=ctx.number("Records", row, "hardbake_time", r.get("hardbake_time")),
        hardbake_time_unit=_time_unit(ctx, row, "hardbake_time_unit", r.get("hardbake_time_unit")),
    )


# ----- public API ---------------------------------------------------------


def load_gold(
    path: Path | str,
    *,
    allow_sets: tuple[str, ...] = ("development",),
    doc_ids: set[str] | None = None,
    split: Split | None = None,
    gold_version: str = "v0",
) -> GoldSet:
    """Load a gold workbook. Documents outside ``allow_sets`` raise ``BenchmarkAccessError``."""
    path = Path(path)
    split = split or get_split()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {name: _rows(wb[name]) for name in wb.sheetnames}
    vocab = get_vocab()

    documents = {d["doc_id"]: d for _, d in sheets.get("Documents", [])}
    present = {r["doc_id"] for _, r in sheets["Records"]}
    for d in sorted(present):
        s = split.set_of(d)
        if s not in allow_sets:
            raise BenchmarkAccessError(
                f"{path.name} contains {d} (set {s!r}); allowed sets: {allow_sets}. "
                "Benchmark annotations are for approved milestone runs only."
            )

    comps: dict[str, list[tuple[int, dict]]] = {}
    for i, c in sheets.get("Components", []):
        comps.setdefault(c["record_id"], []).append((i, c))
    results: dict[str, list[tuple[int, dict]]] = {}
    for i, s in sheets.get("Results", []):
        results.setdefault(s["record_id"], []).append((i, s))
    recon: dict[str, list[dict]] = {}
    for _, x in sheets.get("Reconciliation", []):
        recon.setdefault(x["record_id"], []).append(x)

    out: list[GoldRecord] = []
    for row, r in sheets["Records"]:
        gid, doc_id = str(r["record_id"]), str(r["doc_id"])
        if doc_ids is not None and doc_id not in doc_ids:
            continue
        ctx = _Ctx()
        doc = documents.get(doc_id, {})
        conf = ctx.enum("Records", row, "labeler_confidence", r.get("labeler_confidence"),
                        set(vocab.enum("labeler_confidence")))  # fmt: skip
        comp_rows = sorted(comps.get(gid, []), key=lambda t: (t[1].get("component_no") or 0, t[0]))
        components = [c for i, row_c in comp_rows if (c := _component(ctx, i, row_c))]
        res = [x for i, row_s in results.get(gid, []) if (x := _result(ctx, i, row_s))]
        errors: list[str] = []
        record = None
        try:
            year = ctx.number("Documents", 0, "year", doc.get("year"))
            record = FormulationRecord(
                record_id=gid,
                solids_content_wt_pct=ctx.number(
                    "Records", row, "solids_content_wt_pct", r.get("solids_content_wt_pct")
                ),  # fmt: skip
                provenance=Provenance(
                    source_type=_str(doc.get("source_type")) or "patent",
                    doc_id=doc_id,
                    title=_str(doc.get("title")),
                    assignee_or_authors=_str(doc.get("applicant_or_inventor")),
                    year=int(year) if year else None,
                    location=_str(r.get("location")),
                    evidence_snippet=_str(r.get("evidence_snippet")),
                    is_comparative=_bool(r.get("is_comparative")),
                    extraction_model=None,
                    human_verified=False,  # provisional v0
                    labeler_confidence=conf,
                    labeler_note=_str(r.get("labeler_note")),
                ),
                components=components,
                process=_process(ctx, row, r),
                results=res,
                vocab_version=VOCAB_VERSION,
            )
        except ValidationError as exc:
            errors.append(str(exc))
        out.append(
            GoldRecord(
                gold_record_id=gid,
                doc_id=doc_id,
                example_label=str(r.get("example_label") or ""),
                set=split.set_of(doc_id),
                labeler=_str(r.get("labeler")),
                labeler_confidence=conf,
                low_confidence=conf == "low",
                record=record,
                flags=ctx.flags,
                load_errors=errors,
                original={
                    "record": r,
                    "components": [c for _, c in comp_rows],
                    "results": [s for _, s in results.get(gid, [])],
                },
                reconciliation=recon.get(gid, []),
            )
        )
    return GoldSet(
        path=str(path),
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        gold_version=gold_version,
        documents=[d for k, d in documents.items() if doc_ids is None or k in doc_ids],
        records=out,
    )


def load_development_gold(**kw: Any) -> GoldSet:
    from negative_lines.config import get_settings

    return load_gold(get_settings().gold_dev_path, allow_sets=("development",), **kw)
