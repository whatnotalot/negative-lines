"""Build a document subset of a gold workbook (e.g. the development set) and verify it.

Keeps every sheet, header, style and column of the source workbook and drops
only the rows that belong to other documents:

- Documents, Records, QC_and_Issues, Summary: by ``doc_id``
- Components, Results, Reconciliation: by ``record_id`` -> doc via Records
- Summary ``TOTAL`` rows are dropped (they aggregate all documents)
- README and Lists are kept unchanged; a ``Subset_Info`` sheet records provenance

``verify_subset`` scans every cell (and the package metadata) for identifiers
of forbidden documents, so a subset is safe to commit.

CLI::

    python -m negative_lines.gold.subset --src <full.xlsx> --set development \
        --dst data/gold/v0/development.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import yaml

from negative_lines.gold.split import doc_id_patterns, load_split

DOC_SHEETS = ("Documents", "Records", "QC_and_Issues", "Summary")
RECORD_SHEETS = ("Components", "Results", "Reconciliation")
KEEP_SHEETS = ("README", "Lists")


def sha256(path: Path | str) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


@dataclass
class SubsetReport:
    kept: dict[str, int] = field(default_factory=dict)
    dropped: dict[str, int] = field(default_factory=dict)
    unknown_sheets: list[str] = field(default_factory=list)


def _header(ws) -> list[str]:
    return [c.value for c in ws[1]]


def subset_workbook(
    src: Path | str, dst: Path | str, doc_ids: set[str], *, set_name: str, split_version: int
) -> SubsetReport:
    src, dst = Path(src), Path(dst)
    wb = openpyxl.load_workbook(src)
    report = SubsetReport()

    rec_ws = wb["Records"]
    h = _header(rec_ws)
    rid_i, did_i = h.index("record_id"), h.index("doc_id")
    doc_of = {r[rid_i]: r[did_i] for r in rec_ws.iter_rows(min_row=2, values_only=True) if r[rid_i]}

    for ws in wb.worksheets:
        if ws.title in KEEP_SHEETS:
            continue
        if ws.title in DOC_SHEETS:
            key = _header(ws).index("doc_id")

            def keep(row, key=key):
                return row[key] in doc_ids
        elif ws.title in RECORD_SHEETS:
            key = _header(ws).index("record_id")

            def keep(row, key=key):
                return doc_of.get(row[key]) in doc_ids
        else:
            report.unknown_sheets.append(ws.title)
            continue
        kept = dropped = 0
        for idx in range(ws.max_row, 1, -1):  # bottom-up so indices stay valid
            row = [c.value for c in ws[idx]]
            if all(v is None for v in row):
                continue
            if keep(row):
                kept += 1
            else:
                ws.delete_rows(idx)
                dropped += 1
        report.kept[ws.title], report.dropped[ws.title] = kept, dropped

    if report.unknown_sheets:
        raise ValueError(
            f"unknown sheets, refusing to guess how to subset: {report.unknown_sheets}"
        )

    info = wb.create_sheet("Subset_Info")
    rows = [
        ("field", "value"),
        ("subset_of", src.name),
        ("source_sha256", sha256(src)),
        ("set", set_name),
        ("split_version", split_version),
        ("doc_ids", ", ".join(sorted(doc_ids))),
        ("created_utc", datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")),
        (
            "note",
            "README and Lists sheets are copied unchanged from the full workbook; their "
            "totals describe the full set. Summary TOTAL row removed.",
        ),  # fmt: skip
    ]
    for r in rows:
        info.append(r)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return report


def verify_subset(path: Path | str, allowed: set[str], forbidden: set[str]) -> list[str]:
    """Problems that make the subset unsafe or inconsistent (empty list = OK)."""
    path = Path(path)
    problems: list[str] = []
    patterns = {d: doc_id_patterns(d) for d in forbidden}
    wb = openpyxl.load_workbook(path, read_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                for doc, pats in patterns.items():
                    if any(p in c.value for p in pats):
                        problems.append(f"{ws.title}!{c.coordinate} mentions {doc}")
    # package metadata (properties, sheet names)
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                continue
            text = z.read(name).decode("utf-8", errors="ignore")
            for doc, pats in patterns.items():
                if any(p in text for p in pats):
                    problems.append(f"{name} mentions {doc}")
    # referential integrity and allowed documents
    rec = wb["Records"]
    h = _header(rec)
    doc_of = {
        r[h.index("record_id")]: r[h.index("doc_id")]
        for r in rec.iter_rows(min_row=2, values_only=True)
        if r[0]
    }
    for d in set(doc_of.values()) - allowed:
        problems.append(f"Records contains non-allowed doc {d}")
    for name in RECORD_SHEETS:
        ws = wb[name]
        k = _header(ws).index("record_id")
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[k] and r[k] not in doc_of:
                problems.append(f"{name}: record_id {r[k]} not in Records")
    for name in ("Documents", "QC_and_Issues", "Summary"):
        ws = wb[name]
        k = _header(ws).index("doc_id")
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[k] and r[k] not in allowed:
                problems.append(f"{name}: doc {r[k]} not allowed")
    return problems


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--src", required=True)
    ap.add_argument("--dst", required=True)
    ap.add_argument("--set", default="development")
    ap.add_argument("--split", default=None, help="split.yaml (default: eval/split.yaml)")
    args = ap.parse_args(argv)

    split = load_split(args.split) if args.split else load_split()
    allowed = set(split.doc_ids(args.set))
    forbidden = {d.doc_id for d in split.docs} - allowed
    report = subset_workbook(
        args.src, args.dst, allowed, set_name=args.set, split_version=split.raw["split_version"]
    )
    problems = verify_subset(args.dst, allowed, forbidden)
    meta = {
        "set": args.set,
        "gold_version": split.gold_version,
        "split_version": split.raw["split_version"],
        "source_filename": Path(args.src).name,
        "source_sha256": sha256(args.src),
        "subset_sha256": sha256(args.dst),
        "doc_ids": sorted(allowed),
        "rows_kept": report.kept,
        "verification_problems": problems,
    }
    meta_path = Path(args.dst).with_suffix(".meta.yaml")
    meta_path.write_text(
        yaml.safe_dump(meta, sort_keys=False, allow_unicode=True), encoding="utf-8"
    )
    print(
        yaml.safe_dump(
            {k: v for k, v in meta.items() if k != "verification_problems"}, sort_keys=False
        )
    )
    if problems:
        print(f"VERIFICATION FAILED ({len(problems)} problems):")
        for p in problems[:50]:
            print("  -", p)
        return 1
    print("verification OK: no forbidden document identifiers, references consistent")
    return 0


if __name__ == "__main__":
    sys.exit(main())
