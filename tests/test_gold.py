"""Gold v0: split definition, development subset, adapter and v0->v3 mapping.

Synthetic workbooks are built in tmp_path (no benchmark content anywhere).
The committed development workbook is used for offline smoke tests.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
import yaml

from negative_lines.config import REPO_ROOT, get_settings
from negative_lines.gold.adapter import BenchmarkAccessError, load_development_gold, load_gold
from negative_lines.gold.split import doc_id_patterns, load_split
from negative_lines.gold.subset import sha256, subset_workbook, verify_subset
from negative_lines.gold.vocab_map import (
    normalize_dose_unit,
    normalize_exposure_source,
    normalize_metric,
    parse_mixture,
)

SPLIT = load_split()
DEV = set(SPLIT.doc_ids("development"))
BENCH = set(SPLIT.doc_ids("benchmark"))

REC_COLS = [
    "record_id", "doc_id", "example_label", "is_comparative", "location", "evidence_snippet",
    "solids_content_wt_pct", "substrate", "primer_or_underlayer", "coating_method",
    "film_thickness", "film_thickness_unit", "thickness_stage", "softbake_temp", "softbake_time",
    "softbake_time_unit", "exposure_source", "exposure_mode", "exposure_tool", "dose", "dose_unit",
    "dose_type", "intensity_mw_cm2", "exposure_time_s", "ebeam_kv", "peb_temp", "peb_time",
    "peb_time_unit", "developer_name", "developer_concentration", "developer_time",
    "developer_time_unit", "developer_method", "rinse", "hardbake_temp", "hardbake_time",
    "hardbake_time_unit", "labeler", "labeler_confidence", "labeler_note",
]  # fmt: skip
COMP_COLS = [
    "record_id", "component_no", "role", "role_other_text", "name_as_written",
    "definition_location", "polymer_repeat_units", "polymer_monomer_ratio", "polymer_mw",
    "polymer_pdi", "amount_raw", "amount_unit", "amount_basis", "labeler_confidence",
    "labeler_note",
]  # fmt: skip
RES_COLS = [
    "record_id", "metric", "metric_as_written", "value", "unit", "value_qualitative", "scale",
    "pattern_type", "pitch_nm", "film_thickness", "condition", "method", "sigma",
    "reference_record_id", "labeler_confidence", "labeler_note",
]  # fmt: skip
DOC_COLS = ["file", "doc_id", "source_type", "title", "applicant_or_inventor", "year",
            "pages_read", "records", "notes"]  # fmt: skip
RECON_COLS = ["record_id", "sheet", "field", "labeler_A_value", "labeler_B_value", "decision",
              "agreed_value", "cause", "discrepancy_item_id", "note"]  # fmt: skip


def _sheet(wb, name, cols, rows):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])


def synthetic_workbook(path: Path, doc_id: str = "SYNTH-DEV-1") -> Path:
    """Gold-v0-shaped workbook with the known v0 quirks, for a synthetic document."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "README", ["Project Negative Lines (synthetic)"], [])
    _sheet(wb, "Documents", DOC_COLS, [{"file": "x.pdf", "doc_id": doc_id, "source_type": "patent",
           "title": "Synthetic", "applicant_or_inventor": "Example Co.", "year": 2025}])  # fmt: skip
    base = {"doc_id": doc_id, "location": "[0050]; Table 1", "evidence_snippet": "Example 1 ...",
            "labeler": "Claude (AI), single pass"}  # fmt: skip
    _sheet(wb, "Records", REC_COLS, [
        base | {"record_id": f"{doc_id}_ex1", "example_label": "Example 1", "is_comparative": False,
                "primer_or_underlayer": "HMDS surface modification treatment",
                "softbake_temp": 95, "softbake_time": 2, "softbake_time_unit": "min",
                "exposure_source": "other", "exposure_tool": "aligner; i-line and g-line mixed line",
                "dose": "30-70", "dose_unit": "mJ", "film_thickness": 1.5,
                "film_thickness_unit": "µm", "developer_time": 60, "developer_time_unit": "s",
                "labeler_confidence": "high"},
        base | {"record_id": f"{doc_id}_cex1", "example_label": "Comparative Example 1",
                "is_comparative": 1, "primer_or_underlayer": "lower organic layer: novolak",
                "exposure_source": "ebeam", "dose": 30, "dose_unit": "uC/cm2", "ebeam_kv": 50,
                "labeler_confidence": "low"},
    ])  # fmt: skip
    _sheet(wb, "Components", COMP_COLS, [
        {"record_id": f"{doc_id}_ex1", "component_no": 2, "role": "pag", "name_as_written": "PAG3",
         "amount_raw": "10 parts total; PAG3:PAG4 = 95:5", "amount_unit": "parts",
         "amount_basis": "parts_per_100_resin", "labeler_confidence": "high"},
        {"record_id": f"{doc_id}_ex1", "component_no": 1, "role": "resin", "name_as_written": "R1",
         "amount_raw": 100, "amount_unit": "parts", "amount_basis": "parts_per_100_resin",
         "polymer_mw": 8000, "polymer_pdi": 1.6, "labeler_confidence": "high"},
        {"record_id": f"{doc_id}_ex1", "component_no": 3, "role": "solvent",
         "name_as_written": "PGMEA", "amount_raw": "residual quantity", "amount_unit": "%",
         "amount_basis": "unspecified", "labeler_confidence": "medium"},
        {"record_id": f"{doc_id}_cex1", "component_no": 1, "role": "resin",
         "name_as_written": "R1", "amount_raw": 100, "amount_basis": "parts_per_100_resin"},
    ])  # fmt: skip
    _sheet(wb, "Results", RES_COLS, [
        {"record_id": f"{doc_id}_ex1", "metric": "other", "metric_as_written": "viscosity",
         "value": 25, "unit": "cps"},
        {"record_id": f"{doc_id}_ex1", "metric": "other", "metric_as_written": "Peeling",
         "value_qualitative": "○"},
        {"record_id": f"{doc_id}_ex1", "metric": "other", "metric_as_written": "% via cracking",
         "value": 0, "unit": "%"},
        {"record_id": f"{doc_id}_ex1", "metric": "resolution_min_cd_nm",
         "metric_as_written": "resolution", "value": 80, "unit": "nm", "film_thickness": 100},
        {"record_id": f"{doc_id}_cex1", "metric": "not_a_metric", "metric_as_written": "x",
         "value": "n/a"},
    ])  # fmt: skip
    _sheet(wb, "Reconciliation", RECON_COLS, [
        {"record_id": f"{doc_id}_ex1", "sheet": "Records", "field": "dose", "cause": "unit",
         "decision": "agree with A"},
    ])  # fmt: skip
    _sheet(wb, "Lists", ["role"], [])
    wb.save(path)
    return path


def synthetic_split(path: Path, dev=("SYNTH-DEV-1",), bench=("SYNTH-BENCH-1",)) -> Path:
    data = {
        "split_version": 99, "gold_version": "v0",
        "sets": {
            "development": {"documents": [{"num": "1", "doc_id": d, "records": 2} for d in dev]},
            "benchmark": {"documents": [{"num": "2", "doc_id": d, "records": 2} for d in bench]},
        },
    }  # fmt: skip
    path.write_text(yaml.safe_dump(data), encoding="utf-8")
    return path


# ----- split.yaml ---------------------------------------------------------


def test_split_is_document_level_and_disjoint():
    assert DEV == {"US20210018838A1", "WO2004107052A1", "CN120370628A", "US20080241753A1"}
    assert len(BENCH) == 6 and not DEV & BENCH
    assert sum(d.records for d in SPLIT.docs) == 239
    assert sum(d.records for d in SPLIT.docs if d.set == "development") == 127


def test_doc_id_patterns_cover_kind_code_free_ids():
    assert doc_id_patterns("US7432034B2") == ["US7432034B2", "US7432034"]
    assert doc_id_patterns("KR20240036456A") == ["KR20240036456A", "KR20240036456"]


def test_readme_documents_the_split_consistently():
    readme = (REPO_ROOT / "README.md").read_text(encoding="utf-8")
    assert "eval/split.yaml" in readme
    for d in SPLIT.docs:
        assert d.doc_id in readme, f"README does not list {d.doc_id}"


# ----- leakage guards -----------------------------------------------------


def test_no_benchmark_identifiers_in_prompts_or_fixtures():
    roots = [REPO_ROOT / "src" / "negative_lines" / "extract", REPO_ROOT / "tests" / "fixtures"]
    pats = [p for d in BENCH for p in doc_id_patterns(d)]
    for root in roots:
        for f in root.rglob("*"):
            if f.is_file() and f.suffix in {".py", ".txt", ".json", ".html", ".md", ".yaml"}:
                text = f.read_text(encoding="utf-8", errors="ignore")
                hits = [p for p in pats if p in text]
                assert not hits, f"{f} mentions benchmark document(s) {hits}"


def test_no_benchmark_workbook_inside_repo():
    for f in REPO_ROOT.rglob("*.xlsx"):
        if ".venv" in f.parts:
            continue
        assert f.resolve() == get_settings().gold_dev_path.resolve(), f"unexpected workbook {f}"


def test_committed_development_workbook_is_clean():
    path = get_settings().gold_dev_path
    meta = yaml.safe_load(path.with_suffix(".meta.yaml").read_text(encoding="utf-8"))
    assert meta["subset_sha256"] == sha256(path)
    assert meta["verification_problems"] == []
    assert verify_subset(path, DEV, BENCH) == []
    versions = yaml.safe_load((REPO_ROOT / "eval" / "gold_versions.yaml").read_text("utf-8"))
    assert versions["versions"]["v0"]["derived"]["development"]["sha256"] == sha256(path)


# ----- subset tool --------------------------------------------------------


def test_subset_keeps_only_allowed_docs(tmp_path):
    full = synthetic_workbook(tmp_path / "a.xlsx", "SYNTH-DEV-1")
    other = synthetic_workbook(tmp_path / "b.xlsx", "SYNTH-BENCH-1")
    # merge: append rows of b into a
    wa, wb_ = openpyxl.load_workbook(full), openpyxl.load_workbook(other)
    for name in ("Documents", "Records", "Components", "Results", "Reconciliation"):
        for row in wb_[name].iter_rows(min_row=2, values_only=True):
            wa[name].append(row)
    wa.create_sheet("QC_and_Issues").append(["doc_id", "description"])
    wa["QC_and_Issues"].append(["SYNTH-BENCH-1", "secret issue"])
    wa.create_sheet("Summary").append(["doc_id", "Records"])
    wa["Summary"].append(["SYNTH-BENCH-1", 2])
    wa["Summary"].append(["TOTAL / OVERALL", 4])
    wa.save(full)

    dst = tmp_path / "dev.xlsx"
    report = subset_workbook(full, dst, {"SYNTH-DEV-1"}, set_name="development", split_version=1)
    assert report.kept["Records"] == 2 and report.dropped["Records"] == 2
    assert verify_subset(dst, {"SYNTH-DEV-1"}, {"SYNTH-BENCH-1"}) == []
    assert verify_subset(full, {"SYNTH-DEV-1"}, {"SYNTH-BENCH-1"})  # the full one is not clean
    info = dict(openpyxl.load_workbook(dst)["Subset_Info"].iter_rows(min_row=2, values_only=True))
    assert info["source_sha256"] == sha256(full)


# ----- adapter on a synthetic workbook -----------------------------------


@pytest.fixture()
def synth(tmp_path):
    split = load_split(synthetic_split(tmp_path / "split.yaml"))
    return load_gold(synthetic_workbook(tmp_path / "g.xlsx"), split=split)


def test_adapter_builds_records_and_preserves_originals(synth):
    ex1 = synth.get("SYNTH-DEV-1_ex1")
    assert ex1.record is not None and ex1.load_errors == []
    assert ex1.gold_record_id == ex1.record.record_id == "SYNTH-DEV-1_ex1"  # gold ID preserved
    assert ex1.original["record"]["dose"] == "30-70" and ex1.original["record"]["dose_unit"] == "mJ"
    assert [c.name_as_written for c in ex1.record.components] == ["R1", "PAG3", "PGMEA"]
    assert ex1.record.components[0].polymer.mw == 8000.0
    assert ex1.reconciliation[0]["cause"] == "unit"
    p = ex1.record.process
    assert p.primer.startswith("HMDS") and p.underlayer is None
    assert [(s.temp, s.time, s.time_unit) for s in p.softbake_steps] == [(95.0, 2.0, "min")]
    assert p.film_thickness == 1.5 and p.film_thickness_unit == "µm"
    assert ex1.record.provenance.title == "Synthetic" and ex1.record.provenance.year == 2025
    assert ex1.record.provenance.human_verified is False


def test_adapter_flags(synth):
    ex1 = synth.get("SYNTH-DEV-1_ex1")
    flags = {(f.field, f.status) for f in ex1.flags}
    p = ex1.record.process
    # unambiguous normalizations
    assert p.exposure.source == "mixed_lines" and ("exposure_source", "normalized") in flags
    assert ex1.record.results[0].metric == "viscosity" and ("metric", "normalized") in flags
    pag = ex1.record.components[1]
    assert (pag.mixture_id, pag.mixture_ratio) == ("PAG3+PAG4", "95:5")
    assert pag.amount_raw == "10 parts total; PAG3:PAG4 = 95:5"  # original kept
    # ambiguous / invalid: left unset or unchanged, but flagged
    assert ex1.record.results[1].metric == "other" and ("metric", "ambiguous") in flags
    assert p.exposure.dose is None and p.exposure.dose_unit is None
    assert ("dose_unit", "ambiguous") in flags and ("dose", "invalid") in flags
    assert ("film_thickness", "ambiguous") in flags
    # valid 'other' stays without a flag
    assert ex1.record.results[2].metric == "other"
    assert not any(f.note.startswith("'% via cracking'") for f in ex1.flags)
    # guide rule violation kept as labelled
    solvent = ex1.record.components[2]
    assert solvent.amount_basis == "unspecified" and solvent.labeler_confidence == "medium"
    assert ("labeler_confidence", "rule_violation") in flags
    assert solvent.amount_raw == "residual quantity"


def test_adapter_low_confidence_and_invalid_values(synth):
    cex1 = synth.get("SYNTH-DEV-1_cex1")
    assert cex1.low_confidence and cex1.record.provenance.is_comparative is True
    assert cex1.record.process.underlayer.startswith("lower organic layer")
    assert cex1.record.process.exposure.dose == 30.0
    bad = [(f.field, f.status) for f in cex1.flags]
    assert ("metric", "invalid") in bad and ("value", "invalid") in bad
    assert cex1.record.results[0].metric == "other"


def test_benchmark_documents_need_explicit_permission(tmp_path):
    split = load_split(synthetic_split(tmp_path / "split.yaml"))
    book = synthetic_workbook(tmp_path / "bench.xlsx", "SYNTH-BENCH-1")
    with pytest.raises(BenchmarkAccessError):
        load_gold(book, split=split)
    assert len(load_gold(book, split=split, allow_sets=("benchmark",)).records) == 2


# ----- vocab mapping units -------------------------------------------------


@pytest.mark.parametrize(
    "written,expected,status",
    [
        ("Film thickness/um, Center", "film_thickness_measured", "normalized"),
        ("u %", "thickness_uniformity_pct", "normalized"),
        ("Coating Defects", "coating_defects", "normalized"),
        ("thickness loss of the resist pattern", "thickness_loss", "normalized"),
        ("제막성", "film_forming", "normalized"),
        ("경화막의 내약품성", "chemical_resistance", "normalized"),
        ("Peeling defects", "other", "ambiguous"),
        ("edge roughness", "other", "ambiguous"),
        ("Wafer bow", "other", None),
    ],
)
def test_normalize_metric(written, expected, status):
    metric, flag = normalize_metric("other", written)
    assert metric == expected
    assert (flag.status if flag else None) == status


def test_named_gold_metrics_are_never_remapped():
    assert normalize_metric("adhesion", "Peeling") == ("adhesion", None)


def test_exposure_and_dose_unit_mapping():
    assert (
        normalize_exposure_source("other", "metal halide lamp; visible light >= 480 nm")[0]
        == "visible"
    )
    assert normalize_exposure_source("other", "X-ray irradiation apparatus") == ("other", None)
    assert normalize_exposure_source("other", "UV irradiation device")[1].status == "ambiguous"
    assert normalize_exposure_source("ebeam", "anything") == ("ebeam", None)
    assert normalize_dose_unit("uC/cm2") == ("uC/cm2", None)
    assert normalize_dose_unit("MJ/cm2")[1].status == "invalid"
    assert parse_mixture("12.5 parts total; A:B:C = 1:1:2") == ("A+B+C", "1:1:2")
    assert parse_mixture("10 parts") is None


# ----- committed development gold (offline smoke) -------------------------


def test_development_gold_loads_completely():
    g = load_development_gold()
    assert {r.doc_id for r in g.records} == DEV
    assert len(g.records) == 127 and all(r.record is not None for r in g.records)
    assert all(not r.load_errors for r in g.records)
    by_doc = {d.doc_id: d.records for d in SPLIT.docs if d.set == "development"}
    for doc_id, n in by_doc.items():
        assert len(g.for_doc(doc_id)) == n
    n_comp = sum(len(r.original["components"]) for r in g.records)
    assert sum(len(r.record.components) for r in g.records) == n_comp == 752
    assert {r.doc_id for r in g.records if r.low_confidence} == {"CN120370628A"}
    assert g.get("US20080241753A1_cex1").record.provenance.is_comparative is True
