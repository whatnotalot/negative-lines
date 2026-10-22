import os
import openpyxl
import pandas as pd
import numpy as np

def extract_fujifilm_ebeam_data(excel_path=None, output_csv=None):
    """
    Extracts and standardizes the 14 electron-beam formulation records from US7432034B2.
    All records were processed under identical 50 keV tool conditions on 6-inch Si wafers.
    """
    if excel_path is None:
        excel_path = os.path.join(os.path.dirname(__file__), "..", "data", "Negative_Lines_Reconciled_Gold_Set.xlsx")
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_rec = wb['Records']
    ws_comp = wb['Components']
    ws_res = wb['Results']

    # 1. Gather relevant record metadata (e-beam records only: Ex 1-12, Comp Ex 1-2)
    records = {}
    for r in range(2, ws_rec.max_row + 1):
        rid = ws_rec.cell(r, 1).value
        # Filter for e-beam records of US7432034B2 (Ex 1 to 12, and Comp Ex 1 to 2)
        if rid and str(rid).startswith("US7432034B2") and not str(rid).endswith("ex13") and not str(rid).endswith("cex3") and not str(rid).endswith("cex4"):
            records[rid] = {
                "record_id": rid,
                "example_label": ws_rec.cell(r, 3).value,
                "is_comparative": bool(ws_rec.cell(r, 4).value),
                "exposure_source": "ebeam_50keV",
                "film_thickness_um": 0.30,
                "peb_temp_c": float(ws_rec.cell(r, 26).value or 110.0),
                "peb_time_s": float(ws_rec.cell(r, 27).value or 90.0),
                # Chemical loadings (in grams)
                "resin_amt": 0.0,
                "resin_name": "",
                "mm_crosslinker_amt": 0.0,
                "mm_crosslinker_name": "",
                "cl_crosslinker_amt": 0.0,
                "cl_crosslinker_name": "",
                "pag_amt": 0.0,
                "pag_name": "",
                "quencher_amt": 0.0,
                "quencher_name": "",
                "solvent_pgmea_amt": 7.5,
                "solvent_pgme_amt": 1.5,
                "surfactant_amt": 0.001
            }

    # 2. Extract Components
    for r in range(2, ws_comp.max_row + 1):
        rid = ws_comp.cell(r, 1).value
        if rid in records:
            role = ws_comp.cell(r, 3).value
            name = str(ws_comp.cell(r, 5).value or "").strip()
            amt_raw = ws_comp.cell(r, 11).value
            amt = float(amt_raw) if amt_raw is not None else 0.0

            if role == "resin":
                records[rid]["resin_amt"] = amt
                records[rid]["resin_name"] = name
            elif role == "crosslinker":
                if "MM" in name or name.startswith("MM-"):
                    records[rid]["mm_crosslinker_amt"] = amt
                    records[rid]["mm_crosslinker_name"] = name
                else: # CL series or C-3
                    records[rid]["cl_crosslinker_amt"] = amt
                    records[rid]["cl_crosslinker_name"] = name
            elif role == "pag":
                records[rid]["pag_amt"] = amt
                records[rid]["pag_name"] = name
            elif role == "quencher" or (role == "additive" and "D-" in name):
                records[rid]["quencher_amt"] = amt
                records[rid]["quencher_name"] = name

    # 3. Extract Lithographic Results
    for r in range(2, ws_res.max_row + 1):
        rid = ws_res.cell(r, 1).value
        if rid in records:
            metric = ws_res.cell(r, 2).value
            val = ws_res.cell(r, 4).value
            if val is not None:
                try:
                    num_val = float(str(val).replace("µm", "").replace("nm", "").strip())
                    if metric == "ler_3sigma_nm":
                        records[rid]["ler_3sigma_nm"] = num_val
                    elif metric == "sensitivity_esize":
                        records[rid]["sensitivity_esize_uc_cm2"] = num_val
                    elif metric == "resolution_min_cd_nm":
                        # If printed in um (e.g. 0.11), convert to nm (110)
                        records[rid]["resolution_nm"] = num_val * 1000.0 if num_val < 5.0 else num_val
                except ValueError:
                    pass

    df = pd.DataFrame(list(records.values()))
    
    # Feature engineering
    df["total_crosslinker_amt"] = df["mm_crosslinker_amt"] + df["cl_crosslinker_amt"]
    df["mm_ratio"] = df.apply(
        lambda row: row["mm_crosslinker_amt"] / row["total_crosslinker_amt"] if row["total_crosslinker_amt"] > 0 else 0.0,
        axis=1
    )
    df["has_dual_crosslinker"] = ((df["mm_crosslinker_amt"] > 0) & (df["cl_crosslinker_amt"] > 0)).astype(int)
    df["pag_to_resin_ratio"] = df["pag_amt"] / df["resin_amt"]
    df["quencher_to_pag_ratio"] = df["quencher_amt"] / df["pag_amt"]

    if output_csv is None:
        output_csv = os.path.join(os.path.dirname(__file__), "..", "data", "fujifilm_ebeam_dataset.csv")
    
    df.to_csv(output_csv, index=False)
    return df

def load_ebeam_dataset(csv_path=None):
    """
    Returns feature matrix X, target y_lwr, target y_sens, feature_names, and full dataframe.
    """
    if csv_path is None or not os.path.exists(csv_path):
        default_csv = os.path.join(os.path.dirname(__file__), "..", "data", "fujifilm_ebeam_dataset.csv")
        if not os.path.exists(default_csv):
            extract_fujifilm_ebeam_data(output_csv=default_csv)
        csv_path = default_csv

    df = pd.read_csv(csv_path)

    feature_cols = [
        "resin_amt",
        "mm_crosslinker_amt",
        "cl_crosslinker_amt",
        "pag_amt",
        "quencher_amt",
        "total_crosslinker_amt",
        "mm_ratio",
        "has_dual_crosslinker",
        "pag_to_resin_ratio",
        "quencher_to_pag_ratio"
    ]

    X = df[feature_cols].copy()
    y_lwr = df["ler_3sigma_nm"].values
    y_sens = df["sensitivity_esize_uc_cm2"].values

    return X, y_lwr, y_sens, feature_cols, df

if __name__ == "__main__":
    df = extract_fujifilm_ebeam_data()
    print("Extracted e-beam dataset shape:", df.shape)
    print("\nColumns:", df.columns.tolist())
    print("\nFirst 3 rows:")
    print(df[["record_id", "resin_amt", "mm_crosslinker_amt", "cl_crosslinker_amt", "pag_amt", "quencher_amt", "ler_3sigma_nm", "sensitivity_esize_uc_cm2"]].head(3))
